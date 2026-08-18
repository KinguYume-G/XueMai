"""
测试DOCX文件处理

这个脚本用于测试 python-docx 是否能正常工作
"""
import sys
import os
from io import BytesIO

# 添加项目路径
sys.path.insert(0, r'C:\Users\hp\Desktop\UniPulse Asia\XueMai\backend')

# 设置Django环境
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings.development')
import django
django.setup()

from apps.ai.services.file_processor import FileProcessor

print("=" * 60)
print("测试 DOCX 处理能力")
print("=" * 60)

# 检查 python-docx 是否可用
try:
    import docx
    print(f"✅ python-docx 已安装，版本: {docx.__version__}")
except ImportError as e:
    print(f"❌ python-docx 未安装: {e}")
    sys.exit(1)

# 创建一个简单的测试DOCX文件
try:
    print("\n📝 创建测试 DOCX 文件...")
    test_doc = docx.Document()
    test_doc.add_paragraph("这是测试内容第一段")
    test_doc.add_paragraph("这是测试内容第二段")
    test_doc.add_paragraph("姓名：张三")
    test_doc.add_paragraph("学校：APU")
    
    # 保存到 BytesIO
    buffer = BytesIO()
    test_doc.save(buffer)
    buffer.seek(0)
    
    print("✅ 测试文件创建成功")
    
    # 使用 FileProcessor 处理
    print("\n🔍 使用 FileProcessor 处理...")
    result = FileProcessor.process_file_from_stream(buffer, 'docx')
    
    print("\n📊 处理结果:")
    print(f"  - 提取文本长度: {len(result.get('extracted_text', ''))}")
    print(f"  - 段落数: {result.get('paragraph_count', 0)}")
    print(f"  - 提取文本预览: {result.get('extracted_text', '')[:200]}")
    
    if result.get('extracted_text'):
        print("\n✅ DOCX 处理功能正常")
    else:
        print("\n❌ DOCX 处理失败：提取内容为空")
        
except Exception as e:
    print(f"\n❌ 测试失败: {e}")
    import traceback
    traceback.print_exc()

print("\n" + "=" * 60)
