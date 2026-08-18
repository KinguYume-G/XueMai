"""
文件上传和处理服务

支持的文件类型：
- PDF: 简历、文档、论文
- DOCX: Word文档
- PNG/JPG/JPEG: 图片、截图、课程表
- TXT: 纯文本文件
- CSV: 数据文件
"""

import os
import logging
import mimetypes
from pathlib import Path
from typing import Optional, Dict, Any
import hashlib

# PDF处理
try:
    import pdfplumber
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False
    logging.warning("pdfplumber not installed, PDF support disabled")

# Word文档处理
try:
    import docx
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False
    logging.warning("python-docx not installed, DOCX support disabled")

# 图片OCR处理
try:
    from PIL import Image
    import pytesseract
    OCR_SUPPORT = True
except ImportError:
    OCR_SUPPORT = False
    logging.warning("PIL/pytesseract not installed, OCR support disabled")

logger = logging.getLogger(__name__)


class FileProcessor:
    """文件处理器基类"""
    
    ALLOWED_EXTENSIONS = {
        'pdf': 'application/pdf',
        'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'txt': 'text/plain',
        'png': 'image/png',
        'jpg': 'image/jpeg',
        'jpeg': 'image/jpeg',
        'csv': 'text/csv',
    }
    
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
    
    @classmethod
    def validate_file(cls, file_obj) -> Dict[str, Any]:
        """
        验证文件是否合法
        
        Args:
            file_obj: Django UploadedFile对象
        
        Returns:
            Dict: 验证结果
        
        Raises:
            ValueError: 文件验证失败
        """
        # 检查文件大小
        if file_obj.size > cls.MAX_FILE_SIZE:
            raise ValueError(f"文件大小不能超过 {cls.MAX_FILE_SIZE / 1024 / 1024:.0f}MB")
        
        # 检查文件扩展名
        file_ext = Path(file_obj.name).suffix[1:].lower()
        if file_ext not in cls.ALLOWED_EXTENSIONS:
            raise ValueError(f"不支持的文件类型: {file_ext}")

        if file_ext in {"png", "jpg", "jpeg"} and not OCR_SUPPORT:
            raise ValueError("图片文字识别暂不可用，请上传 PDF、DOCX、TXT 或 CSV")
        
        # 检查MIME类型
        expected_mime = cls.ALLOWED_EXTENSIONS[file_ext]
        actual_mime = file_obj.content_type
        
        if actual_mime and expected_mime != actual_mime:
            raise ValueError(f"文件 MIME 类型不匹配: 期望 {expected_mime}, 实际 {actual_mime}")
        
        return {
            "valid": True,
            "file_name": file_obj.name,
            "file_size": file_obj.size,
            "file_extension": file_ext,
            "mime_type": actual_mime,
        }
    
    @classmethod
    def compute_file_hash(cls, file_obj) -> str:
        """计算文件哈希值（用于去重）"""
        hasher = hashlib.md5()
        for chunk in file_obj.chunks():
            hasher.update(chunk)
        return hasher.hexdigest()
    
    @classmethod
    def process_file(cls, file_path: str, file_type: str) -> Dict[str, Any]:
        """
        处理文件，提取文本内容
        
        Args:
            file_path: 文件路径
            file_type: 文件类型（扩展名）
        
        Returns:
            Dict: 处理结果，包含 extracted_text, metadata等
        """
        if file_type == 'pdf':
            return cls._process_pdf(file_path)
        elif file_type == 'docx':
            return cls._process_docx(file_path)
        elif file_type in ['png', 'jpg', 'jpeg']:
            return cls._process_image(file_path)
        elif file_type == 'txt':
            return cls._process_text(file_path)
        elif file_type == 'csv':
            return cls._process_csv(file_path)
        else:
            raise ValueError(f"不支持的文件类型: {file_type}")
    
    @classmethod
    def process_file_from_stream(cls, file_obj, file_type: str) -> Dict[str, Any]:
        """
        从文件流处理文件（避免路径问题）
        
        Args:
            file_obj: Django UploadedFile对象
            file_type: 文件类型（扩展名）
        
        Returns:
            Dict: 处理结果
        """
        if file_type == 'docx':
            return cls._process_docx_from_stream(file_obj)
        elif file_type == 'pdf':
            return cls._process_pdf_from_stream(file_obj)
        elif file_type == 'txt':
            # TXT直接读取
            content = file_obj.read().decode('utf-8', errors='ignore')
            return {
                "extracted_text": content,
                "line_count": content.count('\n') + 1,
                "metadata": {},
            }
        else:
            # 其他类型仍需要保存到临时文件
            import tempfile
            with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{file_type}') as tmp:
                if hasattr(file_obj, "chunks"):
                    for chunk in file_obj.chunks():
                        tmp.write(chunk)
                else:
                    file_obj.seek(0)
                    tmp.write(file_obj.read())
                tmp_path = tmp.name
            try:
                return cls.process_file(tmp_path, file_type)
            finally:
                os.unlink(tmp_path)
    
    @classmethod
    def _process_pdf(cls, file_path: str) -> Dict[str, Any]:
        """处理PDF文件"""
        if not PDF_SUPPORT:
            raise RuntimeError("PDF处理功能未安装，请安装 pdfplumber")
        
        try:
            with pdfplumber.open(file_path) as pdf:
                text_parts = []
                for page_num, page in enumerate(pdf.pages, 1):
                    text = page.extract_text()
                    if text:
                        text_parts.append(f"[Page {page_num}]\n{text}")
                
                extracted_text = "\n\n".join(text_parts)
                
                return {
                    "extracted_text": extracted_text,
                    "page_count": len(pdf.pages),
                    "metadata": pdf.metadata or {},
                }
        except Exception as e:
            logger.error(f"PDF处理失败: {e}", exc_info=True)
            raise RuntimeError(f"PDF处理失败: {str(e)}")
    
    @classmethod
    def _process_docx(cls, file_path: str) -> Dict[str, Any]:
        """处理Word文档"""
        if not DOCX_SUPPORT:
            raise RuntimeError("DOCX处理功能未安装，请安装 python-docx")
        
        try:
            doc = docx.Document(file_path)
            
            # 提取段落文本
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            # 提取表格文本
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        tables_text.append(row_text)
            
            extracted_text = "\n\n".join(paragraphs)
            if tables_text:
                extracted_text += "\n\n[表格数据]\n" + "\n".join(tables_text)
            
            return {
                "extracted_text": extracted_text,
                "paragraph_count": len(paragraphs),
                "table_count": len(doc.tables),
                "metadata": {},
            }
        except Exception as e:
            logger.error(f"DOCX处理失败: {e}", exc_info=True)
            raise RuntimeError(f"DOCX处理失败: {str(e)}")
    
    @classmethod
    def _process_docx_from_stream(cls, file_obj) -> Dict[str, Any]:
        """从文件流处理Word文档（避免路径问题）"""
        if not DOCX_SUPPORT:
            raise RuntimeError("DOCX处理功能未安装，请安装 python-docx")
        
        try:
            from io import BytesIO
            
            # 读取文件内容到BytesIO
            file_obj.seek(0)
            file_bytes = file_obj.read()
            file_stream = BytesIO(file_bytes)
            
            # 使用BytesIO打开文档
            doc = docx.Document(file_stream)
            
            # 提取段落文本
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            # 提取表格文本
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        tables_text.append(row_text)
            
            # ✅ 新增：提取页眉页脚内容
            headers_footers = []
            for section in doc.sections:
                # 页眉
                if section.header:
                    for para in section.header.paragraphs:
                        if para.text.strip():
                            headers_footers.append(para.text)
                # 页脚
                if section.footer:
                    for para in section.footer.paragraphs:
                        if para.text.strip():
                            headers_footers.append(para.text)
            
            # ✅ 新增：尝试提取文本框和形状中的文字（通过XML）
            textbox_content = []
            try:
                from docx.oxml.text.paragraph import CT_P
                from docx.oxml.table import CT_Tbl
                from docx.oxml import parse_xml
                
                # 遍历文档的XML来查找文本框
                body_element = doc.element.body
                for child in body_element.iter():
                    # 查找文本框元素 (w:txbxContent)
                    if hasattr(child, 'tag') and 'txbxContent' in child.tag:
                        for para_elem in child.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                            para_text = []
                            for text_elem in para_elem.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
                                if text_elem.text:
                                    para_text.append(text_elem.text)
                            if para_text:
                                textbox_content.append(''.join(para_text))
            except Exception as e:
                logger.warning(f"提取文本框内容失败: {e}")
            
            # 组合所有文本
            all_text_parts = []
            
            if paragraphs:
                all_text_parts.append("\n\n".join(paragraphs))
            
            if tables_text:
                all_text_parts.append("\n\n[表格数据]\n" + "\n".join(tables_text))
            
            if headers_footers:
                all_text_parts.append("\n\n[页眉页脚]\n" + "\n".join(headers_footers))
            
            if textbox_content:
                all_text_parts.append("\n\n[文本框内容]\n" + "\n".join(textbox_content))
            
            extracted_text = "\n\n".join(all_text_parts)
            
            logger.info(f"DOCX提取统计: 段落={len(paragraphs)}, 表格={len(doc.tables)}, 页眉页脚={len(headers_footers)}, 文本框={len(textbox_content)}")
            
            return {
                "extracted_text": extracted_text,
                "paragraph_count": len(paragraphs),
                "table_count": len(doc.tables),
                "textbox_count": len(textbox_content),
                "metadata": {},
            }
        except Exception as e:
            logger.error(f"DOCX流处理失败: {e}", exc_info=True)
            raise RuntimeError(f"DOCX流处理失败: {str(e)}")
    
    @classmethod
    def _process_pdf_from_stream(cls, file_obj) -> Dict[str, Any]:
        """从文件流处理PDF"""
        if not PDF_SUPPORT:
            raise RuntimeError("PDF处理功能未安装")
        
        try:
            from io import BytesIO
            file_obj.seek(0)
            pdf_bytes = file_obj.read()
            pdf_stream = BytesIO(pdf_bytes)
            
            with pdfplumber.open(pdf_stream) as pdf:
                text_parts = []
                for page_num, page in enumerate(pdf.pages, 1):
                    text = page.extract_text()
                    if text:
                        text_parts.append(f"[Page {page_num}]\n{text}")
                
                extracted_text = "\n\n".join(text_parts)
                
                return {
                    "extracted_text": extracted_text,
                    "page_count": len(pdf.pages),
                    "metadata": {},
                }
        except Exception as e:
            logger.error(f"PDF流处理失败: {e}", exc_info=True)
            raise RuntimeError(f"PDF流处理失败: {str(e)}")
    
    @classmethod
    def _process_image(cls, file_path: str) -> Dict[str, Any]:
        """处理图片（OCR）"""
        if not OCR_SUPPORT:
            raise RuntimeError("OCR功能未安装，请安装 PIL 和 pytesseract")
        
        try:
            image = Image.open(file_path)
            
            # 执行OCR
            text = pytesseract.image_to_string(image, lang='chi_sim+eng')
            
            return {
                "extracted_text": text.strip(),
                "image_size": image.size,
                "image_format": image.format,
                "metadata": {},
            }
        except Exception as e:
            logger.error(f"图片OCR失败: {e}", exc_info=True)
            # OCR失败不算致命错误，返回空文本
            return {
                "extracted_text": "",
                "error": f"OCR识别失败: {str(e)}",
                "metadata": {},
            }
    
    @classmethod
    def _process_text(cls, file_path: str) -> Dict[str, Any]:
        """处理纯文本文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            
            return {
                "extracted_text": text,
                "line_count": text.count('\n') + 1,
                "metadata": {},
            }
        except UnicodeDecodeError:
            # 尝试其他编码
            try:
                with open(file_path, 'r', encoding='gbk') as f:
                    text = f.read()
                return {
                    "extracted_text": text,
                    "line_count": text.count('\n') + 1,
                    "metadata": {"encoding": "gbk"},
                }
            except Exception as e:
                logger.error(f"文本文件读取失败: {e}", exc_info=True)
                raise RuntimeError(f"文本文件读取失败: {str(e)}")
    
    @classmethod
    def _process_csv(cls, file_path: str) -> Dict[str, Any]:
        """处理CSV文件"""
        try:
            import csv
            
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                rows = list(reader)
            
            # 转换为文本表示
            text_lines = []
            for row in rows:
                text_lines.append(" | ".join(row))
            
            extracted_text = "\n".join(text_lines)
            
            return {
                "extracted_text": extracted_text,
                "row_count": len(rows),
                "column_count": len(rows[0]) if rows else 0,
                "metadata": {},
            }
        except Exception as e:
            logger.error(f"CSV文件处理失败: {e}", exc_info=True)
            raise RuntimeError(f"CSV文件处理失败: {str(e)}")



